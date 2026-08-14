"""
知识库文档导入脚本（通过后端 API）
读取 data/raw/*.docx 文档，通过管理后台 API 上传至知识库

使用方法：
  python data/scripts/import_to_backend.py
  python data/scripts/import_to_backend.py --api-url http://localhost:8000
  python data/scripts/import_to_backend.py --dry-run  # 仅预览不实际上传
"""
import os, sys, json, argparse, re
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("请先安装 python-docx: pip install python-docx")
    sys.exit(1)

try:
    import requests
except ImportError:
    print("请先安装 requests: pip install requests")
    sys.exit(1)

# ---- 配置 ----
RAW_DIR = os.path.join(os.path.dirname(__file__), "..", "raw")
API_BASE = os.environ.get("NEXT_PUBLIC_API_URL", "http://localhost:8000")

# 10 篇文档的元信息映射
DOC_META = {
    "灵山胜境历史文化概述.docx": {
        "category": "历史文化",
        "status": "向量化完成",
        "desc": "全面介绍灵山胜境的历史渊源、五方五佛格局、命名由来、建设历程及文化意义",
    },
    "核心景点讲解词汇总.docx": {
        "category": "景点讲解",
        "status": "向量化完成",
        "desc": "灵山大佛、九龙灌浴、天下第一掌、百子戏弥勒、祥符禅寺、五印坛城、阿育王柱七大景点详细讲解词",
    },
    "游客常见问题FAQ.docx": {
        "category": "FAQ",
        "status": "已解析",
        "desc": "涵盖开放时间、门票、交通、表演、烧香、用餐、拍照、季节等12个常见问题的标准回答",
    },
    "景区服务设施介绍.docx": {
        "category": "便民服务",
        "status": "已解析",
        "desc": "详细介绍游客中心、交通停车、餐饮、购物、无障碍、医疗、智慧导览、卫生间、住宿等服务设施",
    },
    "非遗文化项目介绍.docx": {
        "category": "文化特色",
        "status": "向量化完成",
        "desc": "惠山泥人、无锡精微绣、宜兴紫砂、留青竹刻、锡剧等无锡国家级非遗项目介绍",
    },
    "活动日程安排表.docx": {
        "category": "活动信息",
        "status": "已解析",
        "desc": "九龙灌浴场次、吉祥颂演出、年度节庆活动、特色体验（过堂/抄经/茶道）及旺季提示",
    },
    "灵山传说故事集.docx": {
        "category": "历史文化",
        "status": "已解析",
        "desc": "七则灵山经典传说：玄奘命名、杭恽建寺、赵朴初与大佛、九龙灌浴、佛手与福寿、祥符三桥、梵宫白象",
    },
    "梵宫建筑艺术详解.docx": {
        "category": "文化特色",
        "status": "向量化完成",
        "desc": "灵山梵宫建筑设计与东阳木雕、敦煌壁画、琉璃华藏世界等艺术装置的深度解读",
    },
    "无障碍游览指南.docx": {
        "category": "便民服务",
        "status": "已解析",
        "desc": "门票优惠、轮椅通道与租赁、无障碍卫生间、推荐路线、交通停车、贴心服务的完整无障碍指南",
    },
    "祈福礼佛习俗介绍.docx": {
        "category": "历史文化",
        "status": "已解析",
        "desc": "祈福路线、烧香礼仪、摸佛手抱佛脚、五印坛城祈福、佛教基本礼仪及常用祈福语",
    },
}

# ---- 辅助函数 ----
def extract_text(docx_path):
    """从 .docx 文件提取纯文本内容"""
    doc = Document(docx_path)
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def upload_doc(api_base, title, category, content, status="已解析"):
    """通过 API 上传一篇文档"""
    url = f"{api_base}/api/knowledge"
    payload = {
        "title": title,
        "category": category,
        "content": content,
        "status": status,
    }
    try:
        resp = requests.post(url, json=payload, timeout=30)
        if resp.status_code in (200, 201):
            return True, resp.json()
        else:
            return False, f"HTTP {resp.status_code}: {resp.text[:200]}"
    except requests.exceptions.ConnectionError:
        return False, "连接后端失败，请确认 FastAPI 服务已启动"
    except Exception as e:
        return False, str(e)


def main():
    parser = argparse.ArgumentParser(description="导入知识库文档到后端")
    parser.add_argument("--api-url", default=API_BASE, help=f"后端 API 地址 (默认: {API_BASE})")
    parser.add_argument("--dry-run", action="store_true", help="仅预览，不实际上传")
    parser.add_argument("--rebuild-index", action="store_true", help="上传后触发向量索引重建")
    args = parser.parse_args()

    api_base = args.api_url.rstrip("/")

    # 检查后端连通性
    if not args.dry_run:
        try:
            r = requests.get(f"{api_base}/api/admin/knowledge/stats", timeout=5)
            if r.status_code != 200:
                print(f"警告: 后端返回状态码 {r.status_code}，继续尝试上传...")
        except requests.exceptions.ConnectionError:
            print(f"错误: 无法连接到后端 {api_base}")
            print("请确认 FastAPI 服务已启动，或使用 --api-url 指定正确地址")
            sys.exit(1)

    # 扫描 .docx 文件
    if not os.path.isdir(RAW_DIR):
        print(f"错误: 文档目录不存在 {RAW_DIR}")
        sys.exit(1)

    docx_files = sorted([
        f for f in os.listdir(RAW_DIR)
        if f.endswith('.docx') and not f.startswith('~')
    ])

    if not docx_files:
        print("错误: 未找到任何 .docx 文件")
        sys.exit(1)

    print("=" * 70)
    print(f"  灵境云游 - 知识库文档导入工具")
    print(f"  后端地址: {api_base}")
    if args.dry_run:
        print(f"  *** DRY RUN 模式，不会实际上传 ***")
    print(f"  文档目录: {RAW_DIR}")
    print(f"  找到 {len(docx_files)} 篇文档")
    print("=" * 70)
    print()

    results = {"success": 0, "fail": 0, "skip": 0}
    total_chars = 0

    for idx, filename in enumerate(docx_files, 1):
        filepath = os.path.join(RAW_DIR, filename)
        meta = DOC_META.get(filename, {})
        category = meta.get("category", "其他")
        status = meta.get("status", "已解析")
        desc = meta.get("desc", "")

        # 提取文本
        content = extract_text(filepath)
        char_count = len(content)
        total_chars += char_count

        # 用文件名（去掉 .docx）作为标题
        title = filename.replace(".docx", "")

        print(f"[{idx}/{len(docx_files)}] {title}")
        print(f"  分类: {category}  |  大小: {char_count} 字符  |  状态: {status}")
        if desc:
            print(f"  简介: {desc}")

        if args.dry_run:
            print(f"  [DRY RUN] 跳过上传\n")
            results["skip"] += 1
            continue

        ok, result = upload_doc(api_base, title, category, content, status)
        if ok:
            print(f"  ✓ 上传成功 (ID: {result.get('id', '?')})\n")
            results["success"] += 1
        else:
            print(f"  ✗ 上传失败: {result}\n")
            results["fail"] += 1

    # 总结
    print("=" * 70)
    print(f"  导入完成!")
    print(f"  成功: {results['success']}  |  失败: {results['fail']}  |  跳过: {results['skip']}")
    print(f"  总字符数: {total_chars:,}")
    print("=" * 70)

    # 可选：触发向量索引重建
    if args.rebuild_index and not args.dry_run:
        print("\n触发向量索引重建...")
        try:
            r = requests.post(f"{api_base}/api/ai/rebuild-index", timeout=120)
            if r.status_code == 200:
                print("✓ 向量索引重建完成")
            else:
                print(f"✗ 重建失败: HTTP {r.status_code}")
        except Exception as e:
            print(f"✗ 重建请求失败: {e}")

    # 提示官方资料
    print()
    print("提示: 如还需导入 7 篇官方数据集文档，请在 backend 端运行:")
    print("  python backend/scripts/import_lingshan_data.py")
    print("或在 backend 端运行本脚本的同目录 backend_import.py 直接入库。")


if __name__ == "__main__":
    main()
